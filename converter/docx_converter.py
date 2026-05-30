"""DOCX → HTML：基于 python-docx 的纯 Python 实现。

图片/GIF 保存到 _files/ 目录，HTML 以相对路径引用。"""

from pathlib import Path
from lxml import etree

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from .base import Converter
from .utils import (
    emu_to_px,
    esc,
    guess_image_ext,
    normalize_font,
    pt_to_px,
    rgb_to_hex,
    save_media,
)

# Word 段落对齐 → CSS text-align 映射
_ALIGN_MAP = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}

# XML 命名空间
_NS_W  = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_NS_A  = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_R  = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_M  = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_NS_V  = "urn:schemas-microsoft-com:vml"
_NS_O  = "urn:schemas-microsoft-com:office:office"


class DocxConverter(Converter):
    """将 .docx 文件转换为 HTML，媒体文件保存到 _files/ 目录。"""

    def convert(self, input_path: Path, output_path: Path) -> None:
        self._output_path = output_path
        doc = Document(str(input_path))
        math_xslt = self._load_math_xslt()

        # 图片缓存：{rId: (blob, ext)}
        image_cache = self._build_image_cache(doc)
        video_files = self._build_video_cache(doc, output_path)
        link_map    = self._build_hyperlink_map(doc)
        list_info   = self._build_list_info(doc)

        body_elements = list(doc.element.body)
        # 书签 → 目标段落映射（用于目录内部跳转）
        bm_map = self._build_bookmark_map(body_elements)
        grouped = self._group_list_paragraphs(body_elements, list_info)

        body_parts = []
        for item in grouped:
            if isinstance(item, list):
                body_parts.append(self._render_list_block(
                    item, list_info, doc, image_cache, bm_map, link_map, math_xslt))
            else:
                tag = item.tag.split("}")[-1]
                if tag == "p":
                    body_parts.append(self._render_paragraph(
                        item, doc, image_cache, video_files, bm_map, link_map, math_xslt))
                elif tag == "tbl":
                    body_parts.append(self._render_table(
                        item, doc, image_cache, video_files, bm_map, link_map, math_xslt))
                elif tag == "oMathPara":
                    body_parts.append(self._render_math_block(item, math_xslt))
                elif tag == "hyperlink":
                    body_parts.append(self._render_body_hyperlink(
                        item, doc, image_cache, video_files, bm_map, link_map, math_xslt))
                elif tag == "sdt":
                    body_parts.append(self._render_sdt(
                        item, doc, image_cache, video_files, bm_map, link_map, math_xslt))

        html_body = "\n".join(body_parts)
        full_html = self.render_template("base.html", body=html_body, title=input_path.stem)
        output_path.write_text(full_html, encoding="utf-8")
        if video_files:
            print(f"  媒体文件已保存到: {output_path.parent / output_path.stem}_files/")

    # ── 图片缓存 ──────────────────────────────────────────────────

    def _build_image_cache(self, doc) -> dict:
        """返回 {rId: (blob_bytes, ext)} —— 原始图片数据，不编码 base64。"""
        cache = {}
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                blob = rel.target_part.blob
                cache[rel.rId] = (blob, guess_image_ext(blob))
        return cache

    # ── 图片 run 渲染 ─────────────────────────────────────────────

    def _render_drawing_from_run(self, r_el, image_cache: dict) -> str:
        """从 run 的 <w:drawing> 中提取图片，保存到 _files/ 目录。"""
        drawing_el = r_el.find(f'{{{_NS_W}}}drawing')
        if drawing_el is None:
            return ""

        inline = drawing_el.find(f'{{{_NS_WP}}}inline')
        anchor = drawing_el.find(f'{{{_NS_WP}}}anchor')
        container = inline if inline is not None else anchor
        if container is None:
            return ""

        extent = container.find(f'{{{_NS_WP}}}extent')
        if extent is None:
            return ""
        cx_emu = int(extent.get("cx", "0"))
        cy_emu = int(extent.get("cy", "0"))

        blip = container.find(f'.//{{{_NS_A}}}blip')
        if blip is None:
            return ""
        rId = blip.get(f'{{{_NS_R}}}embed')
        entry = image_cache.get(rId)
        if entry is None:
            return ""

        blob, ext = entry
        rel_path = save_media(self._output_path, blob, ext)

        w_px = emu_to_px(cx_emu)
        h_px = emu_to_px(cy_emu)
        extra = "display:block;margin:0.5em auto;" if anchor is not None else ""
        return (
            f"<img src='{rel_path}' "
            f"width='{w_px:.0f}' height='{h_px:.0f}' "
            f"style='display:inline-block;vertical-align:middle;max-width:100%;{extra}' />"
        )

    # ── 其余方法同前，省略注释（见 git 历史）─────────────────────
    # 注：以下方法保持与上一版本一致，仅 _render_run_full 中的
    #     drawing 调用链改为使用上面的 _render_drawing_from_run。

    @classmethod
    def _load_math_xslt(cls):
        if not hasattr(cls, "_cached_xslt"):
            xsl_path = Path(__file__).resolve().parent.parent / "templates" / "omml2mml.xsl"
            xsl_tree = etree.parse(str(xsl_path))
            cls._cached_xslt = etree.XSLT(xsl_tree)
        return cls._cached_xslt

    def _build_video_cache(self, doc, output_path: Path) -> dict:
        cache = {}
        for rel in doc.part.rels.values():
            reltype = rel.reltype.lower()
            if not ("media" in reltype or "oleobject" in reltype or "package" in reltype):
                continue
            blob = rel.target_part.blob
            ext = self._guess_video_ext(blob, rel.target_ref)
            if ext is None:
                continue
            rel_path = save_media(output_path, blob, ext.lstrip("."))
            cache[rel.rId] = str(rel_path)
        return cache

    def _guess_video_ext(self, blob: bytes, reference: str) -> str | None:
        ref_lower = reference.lower()
        for ve in (".mp4", ".mov", ".avi", ".wmv", ".webm", ".mkv", ".m4v"):
            if ref_lower.endswith(ve):
                return ve
        if blob[:4] == b"\x00\x00\x00\x18ftyp":
            return ".mp4"
        if blob[:4] == b"RIFF" and b"AVI " in blob[8:12]:
            return ".avi"
        if blob[:4] == b"\x1aE\xdf\xa3":
            return ".webm"
        return None

    def _build_hyperlink_map(self, doc) -> dict:
        links = {}
        for rel in doc.part.rels.values():
            if "hyperlink" in rel.reltype.lower() and rel.target_ref:
                links[rel.rId] = rel.target_ref
        return links

    def _build_list_info(self, doc) -> dict:
        info = {}
        numbering_part = doc.part.numbering_part
        if numbering_part is None:
            return info
        numbering_el = numbering_part._element
        abstract_order = {}
        for abn in numbering_el.findall(f'{{{_NS_W}}}abstractNum'):
            abn_id = abn.get(f'{{{_NS_W}}}abstractNumId')
            if abn_id is None:
                continue
            for lvl in abn.findall(f'{{{_NS_W}}}lvl'):
                num_fmt = lvl.find(f'{{{_NS_W}}}numFmt')
                if num_fmt is not None:
                    abstract_order[abn_id] = (num_fmt.get(f'{{{_NS_W}}}val', '') != 'bullet')
                break
        for num in numbering_el.findall(f'{{{_NS_W}}}num'):
            num_id = num.get(f'{{{_NS_W}}}numId')
            if num_id is None:
                continue
            abn_ref = num.find(f'{{{_NS_W}}}abstractNumId')
            if abn_ref is not None:
                abn_id = abn_ref.get(f'{{{_NS_W}}}val')
                if abn_id and abn_id in abstract_order:
                    info[num_id] = {'ordered': abstract_order[abn_id]}
        return info

    def _build_bookmark_map(self, elements: list) -> dict:
        """扫描所有 w:bookmarkStart（可能在 body 级别或 p 内部），
        返回 {目标段落xml元素: 书签名}。"""
        bm = {}
        pending = None
        for el in elements:
            tag = el.tag.split("}")[-1]
            if tag == "bookmarkStart":
                pending = el.get(f'{{{_NS_W}}}name')
            elif tag == "p":
                if pending:
                    bm[el] = pending
                    pending = None
                # 也检查段落内部的 bookmarkStart
                for child in el:
                    if child.tag.split("}")[-1] == "bookmarkStart":
                        inner_name = child.get(f'{{{_NS_W}}}name')
                        if inner_name:
                            bm[el] = inner_name
            elif tag == "bookmarkEnd":
                pending = None
        return bm

    def _group_list_paragraphs(self, elements: list, list_info: dict) -> list:
        result = []; buf = []
        for el in elements:
            tag = el.tag.split("}")[-1]
            if tag == "p" and self._is_list_para(el, list_info) is not None:
                buf.append(el)
            else:
                if buf: result.append(buf); buf = []
                if tag in ("p", "tbl", "oMathPara", "hyperlink", "sdt"):
                    result.append(el)
        if buf: result.append(buf)
        return result

    def _is_list_para(self, p_el, list_info: dict) -> str | None:
        pPr = p_el.find(f'{{{_NS_W}}}pPr')
        if pPr is not None:
            numPr = pPr.find(f'{{{_NS_W}}}numPr')
            if numPr is not None:
                numId_el = numPr.find(f'{{{_NS_W}}}numId')
                if numId_el is not None:
                    num_id = numId_el.get(f'{{{_NS_W}}}val')
                    if num_id and num_id in list_info:
                        return 'ol' if list_info[num_id].get('ordered', True) else 'ul'
                    return 'ul'
        style_name = ""
        pStyle = pPr.find(f'{{{_NS_W}}}pStyle') if pPr is not None else None
        if pStyle is not None:
            style_name = pStyle.get(f'{{{_NS_W}}}val', '').lower()
        if not style_name: return None
        if 'bullet' in style_name or 'unordered' in style_name: return 'ul'
        if 'number' in style_name or 'ordered' in style_name: return 'ol'
        return None

    def _render_list_block(self, items, list_info, doc, image_cache, bm_map, link_map, math_xslt) -> str:
        list_type = self._is_list_para(items[0], list_info) or 'ul'
        tag = list_type
        li_parts = []
        for p_el in items:
            inner = self._render_paragraph_children(
                p_el, doc, image_cache, {}, bm_map, link_map, math_xslt)
            li_parts.append(f"<li>{inner or '&nbsp;'}</li>")
        return f"<{tag} class='docx-list'>\n{''.join(li_parts)}\n</{tag}>"

    def _render_paragraph(self, p_el, doc, image_cache, video_files, bm_map, link_map, math_xslt) -> str:
        para = None
        for p in doc.paragraphs:
            if p._element is p_el: para = p; break
        if para is not None and self._has_paragraph_border(para):
            return "<hr>"
        tag = "p"
        if para is not None and para.style:
            style_name = para.style.name.lower()
            for level in range(1, 7):
                if f"heading {level}" in style_name or f"标题 {level}" in style_name:
                    tag = f"h{level}"; break
        style_parts = self._para_style(para) if para else ""
        inner = self._render_paragraph_children(
            p_el, doc, image_cache, video_files, bm_map, link_map, math_xslt)
        inner = inner or "&nbsp;"
        pb_html = '<hr style="border-style:dashed;margin-top:1.5em">'
        prefix = ""
        suffix = ""
        if para is not None and self._has_page_break(para):
            if tag.startswith("h"):
                prefix = pb_html
            else:
                suffix = pb_html
        bm_id = bm_map.get(p_el, "")
        id_attr = f" id='{bm_id}'" if bm_id else ""
        style_attr = f" style='{style_parts}'" if style_parts else ""
        return f"{prefix}<{tag}{id_attr}{style_attr}>{inner}</{tag}>{suffix}"

    def _render_paragraph_children(self, p_el, doc, image_cache, video_files, bm_map, link_map, math_xslt) -> str:
        parts = []
        for child in p_el:
            tag = child.tag.split("}")[-1]
            if tag == "r":
                parts.append(self._render_run_full(child, image_cache, doc.part))
            elif tag == "oMath":
                parts.append(self._render_math(child, math_xslt))
            elif tag == "hyperlink":
                parts.append(self._render_inline_hyperlink(
                    child, doc, image_cache, video_files, bm_map, link_map, math_xslt))
            elif tag == "object":
                parts.append(self._render_object(child, video_files))
            elif tag == "sdt":
                parts.append(self._render_sdt(
                    child, doc, image_cache, video_files, bm_map, link_map, math_xslt))
        return "".join(parts)

    def _para_style(self, para) -> str:
        if para is None: return ""
        parts = []
        if para.alignment and para.alignment in _ALIGN_MAP:
            parts.append(f"text-align:{_ALIGN_MAP[para.alignment]}")
        pf = para.paragraph_format
        if pf.line_spacing and pf.line_spacing is not None:
            parts.append(f"line-height:{pf.line_spacing:.2f}")
        if pf.space_before:
            parts.append(f"margin-top:{emu_to_px(pf.space_before):.1f}px")
        if pf.space_after:
            parts.append(f"margin-bottom:{emu_to_px(pf.space_after):.1f}px")
        if pf.first_line_indent:
            parts.append(f"text-indent:{emu_to_px(pf.first_line_indent):.1f}px")
        return "; ".join(parts)

    def _has_paragraph_border(self, para) -> bool:
        try:
            pPr = para._element.find(f'{{{_NS_W}}}pPr')
            if pPr is not None:
                pBdr = pPr.find(f'{{{_NS_W}}}pBdr')
                if pBdr is not None:
                    return pBdr.find(f'{{{_NS_W}}}bottom') is not None
        except Exception: pass
        return False

    def _has_page_break(self, para) -> bool:
        for run in para.runs:
            for br in run._element.findall(f'{{{_NS_W}}}br'):
                if br.get(f'{{{_NS_W}}}type') == 'page':
                    return True
        return False

    def _render_run_full(self, r_el, image_cache: dict, part) -> str:
        drawing_html = self._render_drawing_from_run(r_el, image_cache)
        if drawing_html: return drawing_html
        text_parts = []
        for child in r_el:
            tag = child.tag.split("}")[-1]
            if tag == "t":
                t = child.text or ""
                if child.get("{http://www.w3.org/XML/1998/namespace}space") != "preserve":
                    t = t.strip()
                text_parts.append(esc(t))
            elif tag == "br":
                if child.get(f'{{{_NS_W}}}type') != 'page':
                    text_parts.append("<br>")
            elif tag == "tab": text_parts.append("&emsp;")
            elif tag == "sym":
                char_code = child.get(f'{{{_NS_W}}}char')
                if char_code: text_parts.append(f"&#x{char_code};")
        text = "".join(text_parts)
        if not text: return ""
        styles = self._run_styles_from_xml(r_el)
        if styles: return f"<span style='{styles}'>{text}</span>"
        return text

    def _run_styles_from_xml(self, r_el) -> str:
        rPr = r_el.find(f'{{{_NS_W}}}rPr')
        if rPr is None: return ""
        parts = []
        b = rPr.find(f'{{{_NS_W}}}b')
        if b is not None and b.get(f'{{{_NS_W}}}val') != 'false': parts.append("font-weight:bold")
        i = rPr.find(f'{{{_NS_W}}}i')
        if i is not None and i.get(f'{{{_NS_W}}}val') != 'false': parts.append("font-style:italic")
        u = rPr.find(f'{{{_NS_W}}}u')
        if u is not None and u.get(f'{{{_NS_W}}}val') != 'none': parts.append("text-decoration:underline")
        strike = rPr.find(f'{{{_NS_W}}}strike')
        if strike is not None and strike.get(f'{{{_NS_W}}}val') != 'false': parts.append("text-decoration:line-through")
        sz = rPr.find(f'{{{_NS_W}}}sz')
        if sz is not None:
            half_pts = float(sz.get(f'{{{_NS_W}}}val', '0'))
            if half_pts > 0: parts.append(f"font-size:{pt_to_px(half_pts / 2):.1f}px")
        rFonts = rPr.find(f'{{{_NS_W}}}rFonts')
        if rFonts is not None:
            font_name = (rFonts.get(f'{{{_NS_W}}}ascii') or rFonts.get(f'{{{_NS_W}}}hAnsi') or rFonts.get(f'{{{_NS_W}}}eastAsia'))
            if font_name: parts.append(f"font-family:{normalize_font(font_name)}")
        color = rPr.find(f'{{{_NS_W}}}color')
        if color is not None:
            val = color.get(f'{{{_NS_W}}}val')
            if val: parts.append(f"color:#{val}")
        highlight = rPr.find(f'{{{_NS_W}}}highlight')
        if highlight is not None:
            val = highlight.get(f'{{{_NS_W}}}val')
            if val and val != 'none': parts.append(f"background-color:{val}")
        return "; ".join(parts)

    def _render_math(self, omath_el, xslt) -> str:
        return self._omml_to_mathml(omath_el, xslt) or ""

    def _render_math_block(self, omathpara_el, xslt) -> str:
        mathml = self._omml_to_mathml(omathpara_el, xslt)
        return f"<div class='math-block'>{mathml}</div>" if mathml else ""

    def _omml_to_mathml(self, omml_el, xslt) -> str:
        try:
            result = xslt(omml_el)
            raw = str(result)
            if raw.startswith("<?xml"): raw = raw[raw.index("?>") + 2:].lstrip()
            return raw
        except Exception: return ""

    def _render_body_hyperlink(self, hl_el, doc, image_cache, video_files, bm_map, link_map, math_xslt) -> str:
        return self._render_hyperlink_content(hl_el, doc, image_cache, video_files, bm_map, link_map, math_xslt)

    def _render_inline_hyperlink(self, hl_el, doc, image_cache, video_files, bm_map, link_map, math_xslt) -> str:
        return self._render_hyperlink_content(hl_el, doc, image_cache, video_files, bm_map, link_map, math_xslt)

    def _render_hyperlink_content(self, hl_el, doc, image_cache, video_files, bm_map, link_map, math_xslt) -> str:
        # w:anchor = 内部书签（TOC 跳转），优先级高于外部 URL
        anchor = hl_el.get(f'{{{_NS_W}}}anchor')
        if anchor:
            url = "#" + anchor
        else:
            rId = hl_el.get(f'{{{_NS_R}}}id')
            url = link_map.get(rId, "#") if rId else "#"
        inner_parts = []
        for child in hl_el:
            tag = child.tag.split("}")[-1]
            if tag == "r": inner_parts.append(self._render_run_full(child, image_cache, doc.part))
            elif tag == "oMath": inner_parts.append(self._render_math(child, math_xslt))
        inner = "".join(inner_parts) or esc(url)
        # 内部锚点链接（#开头）在当前页跳转，外部 URL 在新标签页打开
        if url.startswith("#"):
            return f"<a href='{esc(url)}'>{inner}</a>"
        return f"<a href='{esc(url)}' target='_blank' rel='noopener noreferrer'>{inner}</a>"

    def _render_object(self, obj_el, video_files: dict) -> str:
        ole = obj_el.find(f'{{{_NS_O}}}OLEObject')
        rId = ole.get(f'{{{_NS_R}}}id') if ole is not None else None
        if rId is None:
            shape = obj_el.find(f'{{{_NS_V}}}shape')
            if shape is not None:
                imagedata = shape.find(f'{{{_NS_V}}}imagedata')
                if imagedata is not None: rId = imagedata.get(f'{{{_NS_R}}}id')
        if rId and rId in video_files:
            src = video_files[rId]
            return (f"<video class='docx-video' controls preload='metadata' "
                    f"style='max-width:100%'><source src='{esc(src)}'>"
                    f"您的浏览器不支持 video 标签。</video>")
        return ""

    def _render_sdt(self, sdt_el, doc, image_cache, video_files, bm_map, link_map, math_xslt) -> str:
        sdt_content = sdt_el.find(f'{{{_NS_W}}}sdtContent')
        if sdt_content is None: return ""
        parts = []
        for child in sdt_content:
            tag = child.tag.split("}")[-1]
            if tag == "p": parts.append(self._render_paragraph(child, doc, image_cache, video_files, bm_map, link_map, math_xslt))
            elif tag == "tbl": parts.append(self._render_table(child, doc, image_cache, video_files, bm_map, link_map, math_xslt))
            elif tag == "r": parts.append(self._render_run_full(child, image_cache, doc.part))
        return "".join(parts)

    def _render_table(self, tbl_el, doc, image_cache, video_files, bm_map, link_map, math_xslt) -> str:
        for table in doc.tables:
            if table._element is tbl_el:
                rows_html = []
                for row in table.rows:
                    cells = "".join(
                        f"<td>{self._render_cell(cell, doc, image_cache, video_files, bm_map, link_map, math_xslt)}</td>"
                        for cell in row.cells)
                    rows_html.append(f"<tr>{cells}</tr>")
                return f"<table class='docx-table'>{''.join(rows_html)}</table>"
        return ""

    def _render_cell(self, cell, doc, image_cache, video_files, bm_map, link_map, math_xslt) -> str:
        parts = []
        for child in cell._element:
            tag = child.tag.split("}")[-1]
            if tag == "p": parts.append(self._render_paragraph(child, doc, image_cache, video_files, bm_map, link_map, math_xslt))
            elif tag == "tbl": parts.append(self._render_table(child, doc, image_cache, video_files, bm_map, link_map, math_xslt))
        return "".join(parts)
